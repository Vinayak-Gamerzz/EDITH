"""Tests for Zenith Studio Presentation Engine, Component Registry, and Master Pipeline."""
import pytest
from pathlib import Path
import pptx

from zenith.presentation.dsl import PresentationDSL, SlideDSL, SlideElement
from zenith.presentation.styles import (
    VISUAL_STYLES,
    infer_visual_style,
    generate_visual_rhythm,
    INTENSITY_CONSTRAINTS,
)
from zenith.presentation.registry import (
    get_presentation_registry,
    PresentationComponentRegistry,
)
from zenith.presentation.renderer import get_renderer, PPTXRenderer
from zenith.presentation.critic import get_critic, VisualQualityCritic
from zenith.presentation.pipeline import get_master_pipeline, MasterPresentationPipeline
from zenith.tools.design_system import (
    compute_relative_luminance,
    compute_contrast_ratio,
    get_wcag_rating,
    auto_adjust_contrast,
    evaluate_palette_contrast,
    get_color_palette_meta,
    THEMES,
)
from zenith.tools.artifact_pipeline import (
    generate_presentation,
    presentation_plan,
    presentation_search_components,
    presentation_critique,
    presentation_inspect_palette,
    edit_presentation,
    preview_presentation,
)
from zenith.memory import store


def test_presentation_dsl_roundtrip():
    """Verify PresentationDSL serialization to dict/json and deserialization from dict."""
    dsl = PresentationDSL(
        title="Autonomous Systems 2026",
        topic="Sovereign Agent Architecture",
        style="cinematic-futuristic",
        theme="cyberpunk_neon",
        intensity=5,
        visual_rhythm=["high-impact", "breathing-room", "data-story", "closing"],
        slides=[
            SlideDSL(
                slide_number=1,
                layout="cinematic-hero",
                style="cinematic-futuristic",
                intensity=5,
                rhythm="high-impact",
                title="Autonomous Systems 2026",
                eyebrow="EXECUTIVE BRIEFING",
                subtitle="The Era of Sovereign Intelligence",
                elements=[
                    SlideElement(type="typography", variant="kinetic-title", text="Autonomous Systems 2026"),
                    SlideElement(type="3d", variant="chrome-orb", asset="chrome-orb"),
                ],
                transition="zoom",
                speaker_notes="Open with authoritative vision on agent sovereignty.",
            ),
            SlideDSL(
                slide_number=2,
                layout="statement",
                style="cinematic-futuristic",
                intensity=2,
                rhythm="breathing-room",
                title="Single Point of Autonomy",
                eyebrow="THESIS",
                subtitle="True autonomy requires zero human friction in routine pipelines.",
                transition="fade",
            ),
        ],
    )

    data = dsl.to_dict()
    assert data["title"] == "Autonomous Systems 2026"
    assert len(data["slides"]) == 2
    assert data["slides"][0]["layout"] == "cinematic-hero"

    restored = PresentationDSL.from_dict(data)
    assert restored.title == dsl.title
    assert restored.style == dsl.style
    assert len(restored.slides) == 2
    assert restored.slides[0].elements[0].variant == "kinetic-title"


def test_visual_styles_and_rhythm():
    """Verify named visual styles, intensity inference, and rhythm alternation."""
    # 1. Check all 5 named styles exist
    for expected in ["cinematic-futuristic", "premium-editorial", "experimental-creative", "playful-youthful", "minimal-premium"]:
        assert expected in VISUAL_STYLES
        style = VISUAL_STYLES[expected]
        assert style.name
        assert style.font_heading
        assert style.font_body

    # 2. Topic-based style inference
    style1, int1 = infer_visual_style("Next-Gen Autonomous AI Agents")
    assert style1 == "cinematic-futuristic"
    assert int1 == 5

    style2, int2 = infer_visual_style("Swiss Clean Architecture and Minimalist Systems")
    assert style2 == "minimal-premium"
    assert int2 == 2

    style3, int3 = infer_visual_style("Boba Tea Cafe Youth Brand Launch")
    assert style3 == "playful-youthful"

    # 3. Visual rhythm alternation
    rhythm = generate_visual_rhythm(num_slides=6)
    assert len(rhythm) == 6
    assert rhythm[0] == "high-impact"
    assert rhythm[-1] == "final-hero"
    # Ensure variety
    assert len(set(rhythm)) >= 4


def test_component_registry_catalog():
    """Verify presentation component library contains layouts, components, effects, typography, charts, and 3D assets."""
    registry = get_presentation_registry()
    assert len(registry.list_layouts()) >= 10
    assert len(registry.list_components()) >= 10

    # Test search by query
    bento = registry.get("bento")
    assert bento is not None
    assert bento.type == "layout"

    charts = registry.search("chart", type_filter="data")
    assert len(charts) >= 3

    effects = registry.search("light", type_filter="effect")
    assert any(e.id == "light-beam" for e in effects)

    # Test compatible styles filter
    cinematic_items = registry.search("", style="cinematic-futuristic", limit=20)
    assert len(cinematic_items) > 0


def test_critic_and_auto_refinement():
    """Verify VisualQualityCritic detects overflow/density fatigue and auto-refines to passing score."""
    critic = get_critic()

    # Create intentionally congested DSL with long title, excessive bullets, and cards
    congested_dsl = PresentationDSL(
        title="Excessively Wordy Deck That Should Trigger Quality Critic Density Warnings And Automatic Truncation",
        topic="Testing Guardrails",
        style="cinematic-futuristic",
        intensity=5,
        slides=[
            SlideDSL(
                slide_number=1,
                layout="bento",
                title="This Is A Way Too Long Slide Title That Definitely Exceeds The Maximum Permitted Characters Allowed In A Single Title Block",
                elements=[
                    SlideElement(
                        type="card",
                        data={
                            "cards": [
                                {"title": f"Card {i}", "points": ["Sample point"]} for i in range(8)  # max is 4
                            ],
                            "bullets": [
                                f"Bullet point {i} with an unnecessarily protracted elaboration describing trivial implementation facets" for i in range(10)
                            ],
                        }
                    )
                ],
            )
        ]
    )

    # Initial critique should detect issues
    initial_critique = critic.critique(congested_dsl)
    assert initial_critique.score < 85

    # Auto-refine
    refined_dsl, final_critique = critic.auto_refine(congested_dsl, min_score=80)
    assert final_critique.score >= 80
    assert len(refined_dsl.slides[0].elements[0].data["cards"]) <= 4
    assert len(refined_dsl.slides[0].elements[0].data["bullets"]) <= 4


@pytest.mark.anyio
async def test_three_distinct_style_presentations(tmp_path: Path):
    """Generate 3 test presentations with completely different visual styles and verify distinct rendering outputs.
    
    Style 1: Cinematic Futuristic (Cyberpunk Neon)
    Style 2: Minimal Premium (Swiss Clean)
    Style 3: Premium Editorial (Editorial Slate)
    """
    store.init_db()

    # ── Presentation 1: Cinematic Futuristic ──────────────────────────────────
    res1 = await generate_presentation(
        title="Autonomous AI Operating Systems",
        topic="Neural Execution Clusters and Decentralized Agents",
        style="cinematic-futuristic",
        theme="cyberpunk_neon",
        subtitle="Foundations of Self-Directing Intelligence",
        format="both",
    )
    assert "Studio Presentation Created" in res1
    assert "Cinematic Futuristic" in res1
    assert "Cyberpunk Neon" in res1
    assert "/presentation/deck_" in res1

    # Extract deck_id
    import re
    m1 = re.search(r"/presentation/(deck_[a-zA-Z0-9_\-]+)", res1)
    assert m1 is not None
    deck1_id = m1.group(1)

    art1 = store.get_artifact(deck1_id)
    assert art1 is not None
    assert art1["theme"] == "cyberpunk_neon"
    pptx1_path = Path(art1["file_path"])
    assert pptx1_path.exists()
    prs1 = pptx.Presentation(str(pptx1_path))
    assert len(prs1.slides) == 6

    # ── Presentation 2: Minimal Premium ───────────────────────────────────────
    res2 = await generate_presentation(
        title="Executive Strategic Capital Audit",
        topic="Minimalist Institutional Allocation Strategy",
        style="minimal-premium",
        theme="swiss_clean",
        subtitle="Pure Structural Clarity and Fiduciary Rigor",
        format="both",
    )
    assert "Studio Presentation Created" in res2
    assert "Minimal Premium" in res2
    assert "Swiss Clean" in res2

    m2 = re.search(r"/presentation/(deck_[a-zA-Z0-9_\-]+)", res2)
    assert m2 is not None
    deck2_id = m2.group(1)

    art2 = store.get_artifact(deck2_id)
    assert art2 is not None
    assert art2["theme"] == "swiss_clean"
    pptx2_path = Path(art2["file_path"])
    assert pptx2_path.exists()
    prs2 = pptx.Presentation(str(pptx2_path))
    assert len(prs2.slides) == 6

    # ── Presentation 3: Premium Editorial ─────────────────────────────────────
    res3 = await generate_presentation(
        title="Global Architecture and Biophilic Design",
        topic="Contemporary Organic Space and Sustainable Synthesis",
        style="premium-editorial",
        theme="editorial_slate",
        subtitle="Spatial Aesthetics for Next-Century Urban Centers",
        format="both",
    )
    assert "Studio Presentation Created" in res3
    assert "Premium Editorial" in res3

    m3 = re.search(r"/presentation/(deck_[a-zA-Z0-9_\-]+)", res3)
    assert m3 is not None
    deck3_id = m3.group(1)

    art3 = store.get_artifact(deck3_id)
    assert art3 is not None
    pptx3_path = Path(art3["file_path"])
    assert pptx3_path.exists()
    prs3 = pptx.Presentation(str(pptx3_path))
    assert len(prs3.slides) == 6

    # Verify visual distinction: different themes, styles, and distinct layout elements
    assert art1["data"]["style"] != art2["data"]["style"]
    assert art1["data"]["theme"] != art2["data"]["theme"]
    assert art2["data"]["style"] != art3["data"]["style"]


@pytest.mark.anyio
async def test_presentation_planning_and_critique_tools():
    """Verify tool_presentation_plan, tool_presentation_search_components, and tool_presentation_critique."""
    # 1. Plan tool
    plan_out = await presentation_plan(
        title="Quantum Mesh Infrastructure",
        topic="Distributed Cryptographic State",
        num_slides=5,
        style="cinematic-futuristic",
    )
    assert "Presentation Plan: Quantum Mesh Infrastructure" in plan_out
    assert "Cinematic Futuristic" in plan_out
    assert "1. **Quantum Mesh Infrastructure" in plan_out

    # 2. Component search tool
    comp_out = await presentation_search_components(query="timeline", type_filter="layout")
    assert "timeline" in comp_out.lower()

    # 3. Critique tool
    critique_out = await presentation_critique()
    assert "Visual Quality Critique" in critique_out
    assert "Quality Score" in critique_out
    assert "Color Palette & Accessibility Audit" in critique_out


def test_color_palette_contrast_calculations():
    """Verify WCAG 2.1 relative luminance, contrast ratios, and theme compliance ratings."""
    # 1. Luminance extremes
    assert compute_relative_luminance("#000000") == pytest.approx(0.0, abs=1e-3)
    assert compute_relative_luminance("#FFFFFF") == pytest.approx(1.0, abs=1e-3)

    # 2. Maximum contrast ratio (Black vs White)
    ratio_max = compute_contrast_ratio("#FFFFFF", "#000000")
    assert ratio_max == pytest.approx(21.0, abs=0.1)

    # 3. Same color has 1.0:1 ratio
    assert compute_contrast_ratio("#121212", "#121212") == 1.0

    # 4. Ratings
    assert get_wcag_rating(7.5) == "AAA"
    assert get_wcag_rating(5.0) == "AA"
    assert get_wcag_rating(3.5) == "AA-Large"
    assert get_wcag_rating(2.0) == "Fail"

    # 5. Verify all primary built-in themes pass WCAG AA/AAA standards
    evaluated_themes = ["editorial_slate", "boba_bash", "swiss_clean", "terracotta_warm", "nordic_navy", "executive_mono", "cyberpunk_neon"]
    for theme_id in evaluated_themes:
        meta = get_color_palette_meta(theme_id)
        assert meta["theme_id"] == theme_id
        contrast = meta["contrast"]
        assert contrast["compliant"] is True
        # Text on background must exceed 4.5:1 (AA minimum for normal text)
        assert contrast["text_on_bg"]["ratio"] >= 4.5
        # Primary themes should achieve AAA (>= 7.0:1)
        assert contrast["text_on_bg"]["ratio"] >= 7.0
        # Text on card surface must also exceed 4.5:1
        assert contrast["text_on_card"]["ratio"] >= 4.5


def test_critic_contrast_awareness_and_repair():
    """Verify VisualQualityCritic detects low contrast backgrounds and auto-repairs them."""
    critic = get_critic()

    # Create slide with near-white background override and white text (terrible contrast < 2.0:1)
    bad_contrast_dsl = PresentationDSL(
        title="Contrast Test Deck",
        topic="Accessibility Verification",
        style="cinematic-futuristic",
        theme="editorial_slate",
        intensity=3,
        slides=[
            SlideDSL(
                slide_number=1,
                layout="statement",
                title="Slide with Intentionally Broken Contrast",
                background_override="#EEEEEE",  # Light grey against white text (#F4F4F6) -> contrast ~ 1.07:1
            ),
            SlideDSL(
                slide_number=2,
                layout="statement",
                title="Slide With Standard Contrast",
            ),
            SlideDSL(
                slide_number=3,
                layout="closing",
                title="Concluding Synthesis",
            ),
        ]
    )

    # 1. Critique must flag contrast violation
    critique = critic.critique(bad_contrast_dsl)
    contrast_issues = [iss for iss in critique.issues if iss["category"] == "contrast"]
    assert len(contrast_issues) >= 1
    assert "low contrast" in contrast_issues[0]["issue"].lower()

    # 2. Auto-refinement must repair the background override to guarantee >= 4.5:1
    refined_dsl, final_critique = critic.auto_refine(bad_contrast_dsl, min_score=85)
    slide1_bg = refined_dsl.slides[0].background_override
    assert slide1_bg is not None
    # Contrast against theme text (#F4F4F6) must now be >= 4.5:1
    repaired_ratio = compute_contrast_ratio("#F4F4F6", slide1_bg)
    assert repaired_ratio >= 4.5


@pytest.mark.anyio
async def test_presentation_inspect_palette_tool():
    """Verify tool_presentation_inspect_palette generates full markdown token tables and WCAG matrix."""
    # 1. Dark canvas theme inspection
    slate_audit = await presentation_inspect_palette("editorial_slate")
    assert "Color Palette Analysis: **Studio Editorial" in slate_audit
    assert "Dark Canvas (Matte Elegance)" in slate_audit
    assert "WCAG AAA" in slate_audit
    assert "Design System Color Tokens" in slate_audit
    assert "WCAG 2.1 Contrast Ratio Matrix" in slate_audit
    assert "Primary Text on Canvas" in slate_audit
    assert "✅ PASS" in slate_audit

    # 2. Light canvas theme inspection
    swiss_audit = await presentation_inspect_palette("swiss_clean")
    assert "Color Palette Analysis: **Swiss Modern Clean**" in swiss_audit
    assert "Light Canvas (Editorial Air)" in swiss_audit
    assert "WCAG AAA" in swiss_audit
    assert "Primary Text on Canvas" in swiss_audit


def test_presentation_pipeline_palette_metadata():
    """Verify PresentationDSL embeds complete palette tokens and contrast metadata."""
    pipeline = get_master_pipeline()
    dsl = pipeline.plan_presentation(
        title="Autonomous Spatial Robotics",
        topic="Robotics & Computer Vision",
        num_slides=4,
        requested_theme="nordic_navy",
    )

    assert dsl.palette_meta is not None
    assert dsl.palette_meta["theme_id"] == "nordic_navy"
    assert "tokens" in dsl.palette_meta
    assert "canvas_background" in dsl.palette_meta["tokens"]
    assert "contrast" in dsl.palette_meta
    assert dsl.palette_meta["contrast"]["compliant"] is True

