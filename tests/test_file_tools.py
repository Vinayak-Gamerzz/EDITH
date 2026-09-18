"""Tests for Presentation (.pptx), Spreadsheet (.xlsx, .csv), and File Analyzer Tools."""
import asyncio
from pathlib import Path
import pytest
import pptx
from pptx.util import Inches
import openpyxl

from zenith.tools.file_processor import read_presentation, read_spreadsheet, analyze_file
from zenith.core.tools import TOOLS


@pytest.mark.anyio
async def test_read_presentation(tmp_path: Path):
    prs = pptx.Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Strategic Roadmap 2026"
    s1.placeholders[1].text = "Vision and Milestone Deliverables"

    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Key Metrics"
    tbl_shape = s2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    tbl = tbl_shape.table
    tbl.cell(0, 0).text = "Metric"
    tbl.cell(0, 1).text = "Target"
    tbl.cell(1, 0).text = "ARR"
    tbl.cell(1, 1).text = "$10M"

    # Add speaker notes
    notes = s2.notes_slide
    notes.notes_text_frame.text = "Discuss ARR timeline."

    deck_path = tmp_path / "roadmap.pptx"
    prs.save(str(deck_path))

    # Verify tool registration
    assert "read_presentation" in TOOLS
    assert "read_spreadsheet" in TOOLS
    assert "analyze_file" in TOOLS

    # Run read_presentation
    result = await read_presentation(str(deck_path))
    assert "Strategic Roadmap 2026" in result
    assert "Key Metrics" in result
    assert "Metric" in result and "Target" in result
    assert "Speaker Notes" in result
    assert "ARR" in result


@pytest.mark.anyio
async def test_read_spreadsheet_xlsx(tmp_path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q1_Budget"
    ws.append(["Department", "Allocated", "Spent"])
    ws.append(["Engineering", 50000, 42000])
    ws.append(["Marketing", 30000, 28000])

    ws2 = wb.create_sheet(title="Notes")
    ws2.append(["Item", "Remark"])
    ws2.append(["Bonus", "Pending"])

    sheet_path = tmp_path / "budget.xlsx"
    wb.save(str(sheet_path))

    # Test reading full workbook summary
    res = await read_spreadsheet(str(sheet_path))
    assert "Q1_Budget" in res
    assert "Engineering" in res
    assert "Allocated" in res

    # Test reading specific sheet
    res_notes = await read_spreadsheet(str(sheet_path), sheet_name="Notes")
    assert "Pending" in res_notes
    assert "Engineering" not in res_notes


@pytest.mark.anyio
async def test_read_spreadsheet_csv(tmp_path: Path):
    csv_path = tmp_path / "users.csv"
    csv_path.write_text("id,username,active\n101,johndoe,true\n102,janedoe,false\n", encoding="utf-8")

    res_csv = await read_spreadsheet(str(csv_path))
    assert "johndoe" in res_csv
    assert "janedoe" in res_csv
    assert "active" in res_csv


@pytest.mark.anyio
async def test_analyze_file_dispatch(tmp_path: Path):
    # PPTX
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Pitch Deck"
    p_path = tmp_path / "pitch.pptx"
    prs.save(str(p_path))

    # CSV
    c_path = tmp_path / "sample.csv"
    c_path.write_text("colA,colB\n1,2\n", encoding="utf-8")

    # JSON
    j_path = tmp_path / "config.json"
    j_path.write_text('{"app": "Zenith", "version": 2.0}', encoding="utf-8")

    res_p = await analyze_file(str(p_path))
    assert "Pitch Deck" in res_p

    res_c = await analyze_file(str(c_path))
    assert "colA" in res_c

    res_j = await analyze_file(str(j_path))
    assert "Zenith" in res_j
